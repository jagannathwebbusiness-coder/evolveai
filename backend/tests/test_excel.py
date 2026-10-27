"""Excel (.xlsx) upload and analysis: real workbooks, hostile containers, bounded limits."""

import io
import zipfile
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.config import Settings
from app.main import create_app
from app.tools.errors import ToolError
from app.tools.excel import parse_xlsx
from app.tools.registry import ToolRegistry
from app.uploads import UploadStore

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(sheets=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in (
        sheets
        or {
            "Sales": [["item", "revenue", "cost"], ["A", 100, 60], ["B", 200, 120], ["C", None, 90]],
            "Other": [["k", "v"], ["x", 1]],
        }
    ).items():
        sheet = wb.create_sheet(name)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def settings(**changes):
    return Settings(_env_file=None, **changes)


def rezip(data: bytes, add=None, drop=None, content_types=None) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            if drop and item.filename == drop:
                continue
            body = source.read(item.filename)
            if content_types and item.filename == "[Content_Types].xml":
                body = content_types(body.decode()).encode()
            target.writestr(item.filename, body)
        for name, body in (add or {}).items():
            target.writestr(name, body)
    return out.getvalue()


def test_parses_every_nonempty_sheet_with_types():
    frames = parse_xlsx(workbook(), settings())
    assert list(frames) == ["Sales", "Other"]
    sales = frames["Sales"]
    assert list(sales.columns) == ["item", "revenue", "cost"] and len(sales) == 3
    assert sales["revenue"].isna().sum() == 1 and sales["revenue"].sum() == 300


def test_formulas_are_never_evaluated():
    wb = Workbook()
    sheet = wb.active
    sheet.append(["a", "b"])
    sheet.append([1, "=1+1"])
    out = io.BytesIO()
    wb.save(out)
    frame = parse_xlsx(out.getvalue(), settings())["Sheet"]
    assert frame["b"].isna().all()  # No cached value exists, and none is computed.


def test_empty_sheets_are_skipped_and_all_empty_is_an_error():
    frames = parse_xlsx(
        workbook({"Empty": [], "Data": [["a"], [1]]}), settings()
    )
    assert list(frames) == ["Data"]
    with pytest.raises(ToolError) as caught:
        parse_xlsx(workbook({"Empty": [], "Header only": [["a", "b"]]}), settings())
    assert caught.value.code == "empty_xlsx"


@pytest.mark.parametrize(
    "rows, code",
    [
        ([["a", "a"], [1, 2]], "invalid_xlsx_header"),
        ([["a", None, "c"], [1, 2, 3]], "invalid_xlsx_header"),
        ([["a", "b"], [1, 2, 3]], "malformed_xlsx"),
        ([["a"], ["x" * 5000]], "malformed_xlsx"),
    ],
)
def test_malformed_sheets_rejected(rows, code):
    data = workbook({"S": rows})
    with pytest.raises(ToolError) as caught:
        parse_xlsx(data, settings())
    assert caught.value.code == code


def test_row_and_column_limits():
    rows = [["a"]] + [[i] for i in range(6)]
    with pytest.raises(ToolError) as caught:
        parse_xlsx(workbook({"S": rows}), settings(csv_max_rows=5))
    assert caught.value.code == "xlsx_row_limit"
    with pytest.raises(ToolError) as caught:
        parse_xlsx(workbook({"S": [["a", "b", "c"], [1, 2, 3]]}), settings(csv_max_columns=2))
    assert caught.value.code == "invalid_xlsx_header"
    sheets = {f"s{i}": [["a"], [1]] for i in range(3)}
    with pytest.raises(ToolError) as caught:
        parse_xlsx(workbook(sheets), settings(xlsx_max_sheets=2))
    assert caught.value.code == "xlsx_too_many_sheets"


@pytest.mark.parametrize(
    "data, code",
    [
        (b"", "invalid_xlsx_size"),
        (b"not a zip at all", "invalid_xlsx"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100, "invalid_xlsx"),  # legacy .xls
        (b"PK\x03\x04garbage", "invalid_xlsx"),
    ],
)
def test_non_xlsx_inputs_rejected(data, code):
    with pytest.raises(ToolError) as caught:
        parse_xlsx(data, settings())
    assert caught.value.code == code


def test_macros_and_wrong_content_types_rejected():
    macro = rezip(workbook(), add={"xl/vbaProject.bin": b"MZ"})
    with pytest.raises(ToolError) as caught:
        parse_xlsx(macro, settings())
    assert caught.value.code == "xlsx_macros"
    swapped = rezip(
        workbook(),
        content_types=lambda text: text.replace("spreadsheetml.sheet.main+xml", "other"),
    )
    with pytest.raises(ToolError):
        parse_xlsx(swapped, settings())
    missing = rezip(workbook(), drop="xl/workbook.xml")
    with pytest.raises(ToolError):
        parse_xlsx(missing, settings())


def test_zip_bomb_is_rejected_before_parsing():
    bomb = rezip(workbook(), add={"xl/bomb.bin": b"0" * (3 * 1024 * 1024)})
    assert len(bomb) < 100_000  # Tiny on disk, huge expanded.
    with pytest.raises(ToolError) as caught:
        parse_xlsx(bomb, settings(xlsx_max_uncompressed_bytes=1024 * 1024))
    assert caught.value.code == "xlsx_too_large"


def test_upload_size_limit_applies():
    with pytest.raises(ToolError) as caught:
        parse_xlsx(workbook(), settings(upload_max_bytes=128))
    assert caught.value.code == "invalid_xlsx_size"


# ---- API + tool ---------------------------------------------------------------------
@pytest.fixture
def api(tmp_path):
    app = create_app(settings(evolve_workspace_dir=tmp_path / "workspace"))
    with TestClient(app) as client:
        yield client


def upload(api, data, name="book.xlsx", content_type=XLSX, conversation=None):
    form = {"conversation_id": conversation} if conversation else {}
    return api.post("/uploads", files={"file": (name, data, content_type)}, data=form)


def test_upload_returns_sheet_metadata_and_stores_xlsx(api, tmp_path):
    response = upload(api, workbook())
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["kind"] == "xlsx" and body["row_count"] == 3
    assert [s["name"] for s in body["sheets"]] == ["Sales", "Other"]
    stored = list((tmp_path / "workspace" / "uploads").glob("*.xlsx"))
    assert len(stored) == 1 and stored[0].stem == body["upload_id"].replace("-", "")


@pytest.mark.parametrize(
    "name, data, status",
    [
        ("book.xlsm", workbook(), 400),
        ("book.xls", workbook(), 400),
        ("../book.xlsx", workbook(), 400),
        ("book.xlsx", b"PK\x03\x04junk", 400),
        ("book.xlsx", rezip(workbook(), add={"xl/vbaProject.bin": b"MZ"}), 400),
    ],
)
def test_bad_uploads_rejected(api, name, data, status):
    assert upload(api, data, name=name).status_code == status


def test_csv_upload_still_works(api):
    response = upload(api, b"a,b\n1,2\n", name="x.csv", content_type="text/csv")
    assert response.status_code == 201 and response.json()["kind"] == "csv"
    assert response.json()["sheets"] == []


def make_registry(tmp_path, data, name="book.xlsx"):
    config = settings(evolve_workspace_dir=tmp_path / "workspace", evolve_data_dir=tmp_path / "data")
    store = UploadStore(config)
    task = uuid4()
    saved = store.save(data, name, task)
    return ToolRegistry(config, store, task, [saved.upload_id], None), saved


def analyze(registry, saved, **extra):
    args = registry.validate(
        "analyze_csv", {"upload_id": str(saved.upload_id), **extra}
    )
    return registry.execute("analyze_csv", args)


def test_analyze_defaults_to_first_sheet_and_selects_by_name(tmp_path):
    registry, saved = make_registry(tmp_path, workbook())
    first = analyze(registry, saved)
    assert first["sheet"] == "Sales" and first["sheets"] == ["Sales", "Other"]
    assert first["missing_values"]["revenue"] == 1
    assert first["numeric_summary"]["revenue"]["mean"] == 150.0
    other = analyze(registry, saved, sheet="Other")
    assert other["columns"] == ["k", "v"]
    with pytest.raises(ToolError) as caught:
        analyze(registry, saved, sheet="Nope")
    assert caught.value.code == "invalid_sheet"


def test_sheet_argument_on_csv_is_rejected(tmp_path):
    registry, saved = make_registry(tmp_path, b"a,b\n1,2\n", name="x.csv")
    assert analyze(registry, saved)["row_count"] == 1
    with pytest.raises(ToolError) as caught:
        analyze(registry, saved, sheet="Sheet1")
    assert caught.value.code == "invalid_sheet"


def test_uploaded_workbook_is_conversation_scoped(tmp_path):
    registry, saved = make_registry(tmp_path, workbook())
    other = ToolRegistry(registry.settings, registry.uploads, uuid4(), [saved.upload_id], None)
    with pytest.raises(ToolError):
        analyze(other, saved)
