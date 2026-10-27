"""Bounded, read-only .xlsx parsing.

Only cached cell values are read (formulas are never evaluated), macro-enabled and
encrypted workbooks are rejected, and zip expansion, sheet, row, column and cell sizes are
capped before any data reaches pandas.
"""

import io
import zipfile
from datetime import date, datetime, time

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..config import Settings
from .errors import ToolError

MAIN_CONTENT_TYPE = "spreadsheetml.sheet.main+xml"
MAX_ZIP_ENTRIES = 500
MAX_CELL_CHARS = 4096


def _check_container(data: bytes, settings: Settings) -> None:
    if not data or len(data) > settings.upload_max_bytes:
        raise ToolError(
            "invalid_xlsx_size", "The workbook is empty or exceeds UPLOAD_MAX_BYTES."
        )
    if not data.startswith(b"PK\x03\x04"):
        raise ToolError(
            "invalid_xlsx",
            "This is not an .xlsx workbook (encrypted, legacy .xls and renamed files are not accepted).",
        )
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise ToolError("invalid_xlsx", "The workbook is corrupt.") from exc
    with archive:
        if len(entries) > MAX_ZIP_ENTRIES:
            raise ToolError("xlsx_too_complex", "The workbook has too many parts.")
        if sum(entry.file_size for entry in entries) > settings.xlsx_max_uncompressed_bytes:
            raise ToolError(
                "xlsx_too_large",
                "The workbook expands beyond XLSX_MAX_UNCOMPRESSED_BYTES; upload a smaller file.",
            )
        names = {entry.filename.lower() for entry in entries}
        if any("vbaproject" in name for name in names):
            raise ToolError(
                "xlsx_macros",
                "Macro-enabled workbooks are rejected. Save a plain .xlsx without macros.",
            )
        if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
            raise ToolError("invalid_xlsx", "This is not a valid .xlsx workbook.")
        content_types = archive.read("[Content_Types].xml").decode("utf-8", "ignore")
        if MAIN_CONTENT_TYPE not in content_types or "macroEnabled" in content_types:
            raise ToolError(
                "invalid_xlsx",
                "Only plain .xlsx workbooks (no macros) are accepted.",
            )


def _cell(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        if len(value) > MAX_CELL_CHARS:
            raise ToolError(
                "malformed_xlsx",
                f"Cells are limited to {MAX_CELL_CHARS} characters.",
            )
        if any(ord(char) < 32 and char not in "\r\n\t" for char in value):
            raise ToolError(
                "malformed_xlsx", "Control characters are not accepted in cells."
            )
        return value
    if isinstance(value, (bool, int, float)) or value is None:
        return value
    return str(value)[:MAX_CELL_CHARS]


def _sheet_frame(worksheet, settings: Settings) -> pd.DataFrame | None:
    rows = worksheet.iter_rows(values_only=True)
    header = None
    for raw in rows:  # Skip leading fully empty rows.
        if any(cell is not None and str(cell).strip() for cell in raw):
            header = list(raw)
            break
    if header is None:
        return None
    while header and (header[-1] is None or not str(header[-1]).strip()):
        header.pop()
    names = [str(name).strip() if name is not None else "" for name in header]
    if (
        not names
        or any(not name or len(name) > 100 for name in names)
        or len(names) > settings.csv_max_columns
        or len(set(names)) != len(names)
    ):
        raise ToolError(
            "invalid_xlsx_header",
            f"Sheet '{worksheet.title}' needs a first row of unique, nonempty column names (max "
            f"{settings.csv_max_columns} columns, 100 characters each).",
        )
    body = []
    for raw in rows:
        values = [_cell(cell) for cell in raw]
        if all(cell is None for cell in values):
            continue
        if any(cell is not None for cell in values[len(names) :]):
            raise ToolError(
                "malformed_xlsx",
                f"Sheet '{worksheet.title}' has data to the right of its header columns.",
            )
        values = (values + [None] * len(names))[: len(names)]
        body.append(values)
        if len(body) > settings.csv_max_rows:
            raise ToolError(
                "xlsx_row_limit",
                f"Sheet '{worksheet.title}' exceeds CSV_MAX_ROWS; upload a smaller file.",
            )
    if not body:
        return None
    return pd.DataFrame(body, columns=names).infer_objects()


def parse_xlsx(data: bytes, settings: Settings) -> dict[str, pd.DataFrame]:
    """Return {sheet name: validated frame} for every non-empty sheet."""
    _check_container(data, settings)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ToolError("invalid_xlsx", "The workbook could not be read.") from exc
    except Exception as exc:  # openpyxl surfaces assorted XML errors on malformed parts.
        raise ToolError("invalid_xlsx", "The workbook could not be parsed.") from exc
    try:
        if len(workbook.sheetnames) > settings.xlsx_max_sheets:
            raise ToolError(
                "xlsx_too_many_sheets",
                f"The workbook has more than {settings.xlsx_max_sheets} sheets.",
            )
        frames: dict[str, pd.DataFrame] = {}
        for name in workbook.sheetnames:
            try:
                frame = _sheet_frame(workbook[name], settings)
            except ToolError:
                raise
            except Exception as exc:
                raise ToolError(
                    "malformed_xlsx", f"Sheet '{name[:60]}' could not be read."
                ) from exc
            if frame is not None:
                frames[name] = frame
    finally:
        workbook.close()
    if not frames:
        raise ToolError(
            "empty_xlsx",
            "The workbook needs a header row and at least one data row on some sheet.",
        )
    return frames
